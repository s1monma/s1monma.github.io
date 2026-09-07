"""
Microsoft Graph API & SharePoint Excel Sync for AI & Crypto Terms Glossary.

Authenticates to Microsoft Graph using an Azure AD (Microsoft Entra ID) app
registration, downloads the latest glossary Excel from SharePoint, and generates
both site/data.json and site/data.js for the frontend.

Usage:
    python sync.py                      # Sync via Microsoft Graph from SharePoint
    python sync.py --local              # Sync from local Excel file (default path)
    python sync.py --local path.xlsx    # Sync from a specific local Excel file
    python sync.py --test-auth          # Test Azure AD / Graph API authentication
    python sync.py --setup              # Print Azure AD IT setup instructions
"""

from __future__ import annotations

import argparse
import base64
import json
import os
import re
import ssl
import sys
import tempfile
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
SITE_DIR = HERE / "site"
DATA_JSON = SITE_DIR / "data.json"
DATA_JS = SITE_DIR / "data.js"
DEFAULT_LOCAL_EXCEL = HERE.parent / "hardcopy_KS Crypto Collective - Crypto Terms Glossary - 04.30.26 copy.xlsx"

HEADER_ROW = 2
DATA_START_ROW = 3
COLS = {
    "num": 1,
    "ref": 2,
    "term": 3,
    "definition": 4,
    "example": 5,
    "context": 6,
    "child_term": 7,
    "source": 8,
    "citation": 9,
    "litigation": 10,
    "links": 11,
}

URL_RE = re.compile(r"https?://[^\s\")]+")
LIT_CASE_RE = re.compile(r"^\[([^\]]+)\]\s*:?\s*(.*)$", re.DOTALL)


# ═══════════════════════════════════════════════════════════════
# ENVIRONMENT & CONFIGURATION
# ═══════════════════════════════════════════════════════════════
def load_dotenv() -> None:
    """Load key-value pairs from .env into os.environ if not already set."""
    candidates = [
        HERE / ".env",
        HERE.parent / ".env",
    ]
    for env_file in candidates:
        if env_file.is_file():
            try:
                for line in env_file.read_text(encoding="utf-8").splitlines():
                    line = line.strip()
                    if not line or line.startswith("#") or "=" not in line:
                        continue
                    key, val = line.split("=", 1)
                    key = key.strip()
                    val = val.strip().strip("'\"")
                    if key and key not in os.environ:
                        os.environ[key] = val
            except Exception as e:
                print(f"[Warning] Could not read {env_file}: {e}")


def print_it_setup_instructions() -> None:
    print(r"""
================================================================================
MICROSOFT GRAPH API & SHAREPOINT SYNC — IT SETUP INSTRUCTIONS
================================================================================

To allow sync.py to download the glossary Excel from SharePoint automatically:

1. Register an App in Microsoft Entra ID (Azure AD):
   a. Sign in to https://portal.azure.com
   b. Go to: Microsoft Entra ID > App registrations > New registration
   c. Name: "Keystone Glossary Sync" (or similar)
   d. Supported account types: "Accounts in this organizational directory only"
   e. Click Register.

2. Create a Client Secret:
   a. Under "Certificates & secrets", select "Client secrets" > "New client secret"
   b. Description: "Glossary Sync Secret"
   c. Expiration: 12 or 24 months (per your security policy)
   d. Copy the Value immediately (this is your AZURE_CLIENT_SECRET).

3. Grant API Permissions:
   a. Under "API permissions", click "Add a permission" > "Microsoft Graph"
   b. Select "Application permissions"
   c. Add:
      - Files.Read.All (or Sites.Read.All)
   d. Click "Grant admin consent for <Your Organization>" (Requires Admin role).

4. Configure .env file:
   Copy .env.example to .env and fill in:
     AZURE_TENANT_ID=<Directory (tenant) ID from App Overview>
     AZURE_CLIENT_ID=<Application (client) ID from App Overview>
     AZURE_CLIENT_SECRET=<Client secret Value from Step 2>
     SHAREPOINT_SITE_URL=https://<tenant>.sharepoint.com/sites/<sitename>
     SHAREPOINT_FILE_PATH=Shared Documents/KS Crypto Collective - Crypto Terms Glossary.xlsx

================================================================================
""")


# ═══════════════════════════════════════════════════════════════
# MICROSOFT GRAPH API CLIENT
# ═══════════════════════════════════════════════════════════════
class GraphSyncClient:
    def __init__(self, tenant_id: str, client_id: str, client_secret: str):
        self.tenant_id = tenant_id.strip()
        self.client_id = client_id.strip()
        self.client_secret = client_secret.strip()
        self.token: str | None = None
        self.ssl_ctx = ssl.create_default_context()

    def get_token(self) -> str:
        """Acquire an application bearer token via OAuth 2.0 Client Credentials."""
        if self.token:
            return self.token

        token_url = f"https://login.microsoftonline.com/{self.tenant_id}/oauth2/v2.0/token"
        payload = urllib.parse.urlencode({
            "client_id": self.client_id,
            "client_secret": self.client_secret,
            "grant_type": "client_credentials",
            "scope": "https://graph.microsoft.com/.default",
        }).encode("utf-8")

        req = urllib.request.Request(
            token_url,
            data=payload,
            headers={
                "Content-Type": "application/x-www-form-urlencoded",
                "User-Agent": "Keystone-GlossarySync/2.0",
            },
            method="POST",
        )

        try:
            with urllib.request.urlopen(req, context=self.ssl_ctx, timeout=30) as resp:
                data = json.loads(resp.read().decode("utf-8"))
                self.token = data.get("access_token")
                if not self.token:
                    raise RuntimeError("No access_token returned by Microsoft identity endpoint.")
                return self.token
        except urllib.error.HTTPError as e:
            err_body = e.read().decode("utf-8", errors="ignore")
            raise RuntimeError(
                f"Azure AD authentication failed (HTTP {e.code}): {err_body}\n"
                "Please verify AZURE_TENANT_ID, AZURE_CLIENT_ID, and AZURE_CLIENT_SECRET."
            ) from e

    def _graph_request(self, url: str) -> dict:
        token = self.get_token()
        req = urllib.request.Request(
            url,
            headers={
                "Authorization": f"Bearer {token}",
                "Accept": "application/json",
                "User-Agent": "Keystone-GlossarySync/2.0",
            },
        )
        with urllib.request.urlopen(req, context=self.ssl_ctx, timeout=30) as resp:
            return json.loads(resp.read().decode("utf-8"))

    def resolve_site_id(self, site_url: str) -> str:
        """Resolve a full SharePoint URL (e.g. https://domain.sharepoint.com/sites/name) to a Graph site-id."""
        parsed = urllib.parse.urlparse(site_url)
        hostname = parsed.netloc
        site_path = parsed.path.rstrip("/")
        if not hostname or not site_path:
            raise ValueError(f"Invalid SHAREPOINT_SITE_URL format: '{site_url}'")

        graph_site_url = f"https://graph.microsoft.com/v1.0/sites/{hostname}:{site_path}"
        data = self._graph_request(graph_site_url)
        site_id = data.get("id")
        if not site_id:
            raise RuntimeError(f"Could not retrieve site ID for {site_url}")
        return site_id

    def download_file(
        self,
        dest_path: Path,
        site_url: str | None = None,
        file_path: str | None = None,
        sharing_url: str | None = None,
    ) -> Path:
        """Download file content from SharePoint and save to dest_path."""
        token = self.get_token()

        download_urls = []

        if sharing_url:
            # Base64url encode sharing URL with u! prefix
            b64 = base64.b64encode(sharing_url.encode("utf-8")).decode("utf-8")
            encoded_token = "u!" + b64.rstrip("=").replace("/", "_").replace("+", "-")
            download_urls.append(f"https://graph.microsoft.com/v1.0/shares/{encoded_token}/driveItem/content")

        elif site_url and file_path:
            print(f"[Graph API] Resolving SharePoint site: {site_url}")
            site_id = self.resolve_site_id(site_url)

            # Strip leading slashes
            clean_path = file_path.strip().lstrip("/")
            encoded_path = urllib.parse.quote(clean_path)

            # Primary attempt: via root document library
            download_urls.append(f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{encoded_path}:/content")

            # Alternate attempt if path had "Shared Documents/" prefix
            if clean_path.lower().startswith("shared documents/"):
                sub_path = clean_path[len("shared documents/"):].strip("/")
                encoded_sub = urllib.parse.quote(sub_path)
                download_urls.append(f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{encoded_sub}:/content")
        else:
            raise ValueError("Either (SHAREPOINT_SITE_URL and SHAREPOINT_FILE_PATH) or SHAREPOINT_SHARING_URL must be specified.")

        last_err = None
        for d_url in download_urls:
            print(f"[Graph API] Requesting file download from Microsoft Graph...")
            req = urllib.request.Request(
                d_url,
                headers={
                    "Authorization": f"Bearer {token}",
                    "User-Agent": "Keystone-GlossarySync/2.0",
                },
            )
            try:
                with urllib.request.urlopen(req, context=self.ssl_ctx, timeout=60) as resp:
                    dest_path.parent.mkdir(parents=True, exist_ok=True)
                    content = resp.read()
                    dest_path.write_bytes(content)
                    print(f"[Graph API] Successfully downloaded Excel file ({len(content) / 1024:.1f} KB) -> {dest_path.name}")
                    return dest_path
            except urllib.error.HTTPError as e:
                last_err = e
                if e.code == 404 and len(download_urls) > 1:
                    continue
                err_msg = e.read().decode("utf-8", errors="ignore")
                raise RuntimeError(
                    f"Graph API download failed (HTTP {e.code}): {err_msg}\n"
                    f"Verify the file path exists on SharePoint and the Azure AD App has Files.Read.All permission."
                ) from e

        if last_err:
            raise RuntimeError(f"Could not locate file on SharePoint (HTTP {last_err.code}).")
        raise RuntimeError("No download URL could be formed.")


# ═══════════════════════════════════════════════════════════════
# EXCEL PARSING & PROCESSING
# ═══════════════════════════════════════════════════════════════
def clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def split_children(raw: str) -> list[str]:
    if not raw:
        return []
    return [c.strip() for c in raw.split(",") if c.strip()]


def slugify(term: str) -> str:
    s = term.lower().strip()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return s.strip("-")


def first_letter_bucket(term: str) -> str:
    if not term:
        return "#"
    first = term.lstrip()[0].upper()
    if first.isalpha():
        return first
    return "#"


def parse_citation(raw: str) -> dict:
    if not raw:
        return {"text": "", "url": ""}
    url_match = URL_RE.search(raw)
    return {
        "text": raw.strip(),
        "url": url_match.group(0) if url_match else "",
    }


def parse_litigation(raw: str) -> dict:
    if not raw:
        return {"case": "", "description": ""}
    m = LIT_CASE_RE.match(raw.strip())
    if m:
        return {"case": m.group(1).strip(), "description": m.group(2).strip()}
    return {"case": "", "description": raw.strip()}


def read_excel(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet_name = "glossary" if "glossary" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    print(f"Reading sheet '{sheet_name}' from {path.name}")

    terms: list[dict] = []
    for row_idx in range(DATA_START_ROW, DATA_START_ROW + 500):
        term = clean(ws.cell(row=row_idx, column=COLS["term"]).value)
        if not term:
            if row_idx > DATA_START_ROW + 5 and not any(
                ws.cell(row=row_idx + i, column=COLS["term"]).value for i in range(5)
            ):
                break
            continue

        num_raw = ws.cell(row=row_idx, column=COLS["num"]).value
        try:
            num = int(num_raw) if num_raw is not None else len(terms) + 1
        except (TypeError, ValueError):
            num = len(terms) + 1

        entry = {
            "num": num,
            "ref": clean(ws.cell(row=row_idx, column=COLS["ref"]).value),
            "term": term,
            "slug": slugify(term),
            "letter": first_letter_bucket(term),
            "definition": clean(ws.cell(row=row_idx, column=COLS["definition"]).value),
            "example": clean(ws.cell(row=row_idx, column=COLS["example"]).value),
            "context": clean(ws.cell(row=row_idx, column=COLS["context"]).value),
            "related": split_children(clean(ws.cell(row=row_idx, column=COLS["child_term"]).value)),
            "source": clean(ws.cell(row=row_idx, column=COLS["source"]).value),
            "citation": parse_citation(clean(ws.cell(row=row_idx, column=COLS["citation"]).value)),
            "litigation": parse_litigation(clean(ws.cell(row=row_idx, column=COLS["litigation"]).value)),
            "links": clean(ws.cell(row=row_idx, column=COLS["links"]).value),
        }
        terms.append(entry)

    return terms


def enrich_with_existing_metadata(terms: list[dict], existing_json_path: Path) -> None:
    """Preserve research papers, cluster tags, and hierarchy rings from existing data.json."""
    if not existing_json_path.is_file():
        return

    try:
        data = json.loads(existing_json_path.read_text(encoding="utf-8"))
        existing_terms = data.get("terms", [])
        by_slug = {t.get("slug"): t for t in existing_terms if t.get("slug")}
        by_name = {t.get("term", "").lower().strip(): t for t in existing_terms if t.get("term")}

        merged_count = 0
        for t in terms:
            match = by_slug.get(t["slug"]) or by_name.get(t["term"].lower().strip())
            if match:
                if "papers" in match and match["papers"]:
                    t["papers"] = match["papers"]
                if "cluster" in match and match["cluster"]:
                    t["cluster"] = match["cluster"]
                if "ring" in match:
                    t["ring"] = match["ring"]
                merged_count += 1
            else:
                t.setdefault("papers", [])

        print(f"Preserved research papers & cluster tags for {merged_count} existing terms.")
    except Exception as e:
        print(f"[Warning] Could not merge existing metadata from {existing_json_path.name}: {e}")


def build_payload(terms: list[dict]) -> dict:
    total = len(terms)
    with_citation = sum(1 for t in terms if t.get("citation", {}).get("text"))
    with_litigation = sum(
        1 for t in terms if t.get("litigation", {}).get("case") or t.get("litigation", {}).get("description")
    )
    with_papers = sum(1 for t in terms if t.get("papers"))
    unique_clusters = len({t.get("cluster") for t in terms if t.get("cluster")})

    return {
        "meta": {
            "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "total_terms": total,
            "with_citations": with_citation,
            "with_papers": with_papers,
            "litigation_examples": with_litigation,
            "total_clusters": unique_clusters if unique_clusters > 0 else 8,
            "last_updated_display": datetime.now().strftime("%b %d, %Y"),
        },
        "terms": terms,
    }


def write_site_assets(payload: dict) -> None:
    SITE_DIR.mkdir(parents=True, exist_ok=True)

    # 1. Write data.json
    DATA_JSON.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")

    # 2. Write data.js (enables zero-CORS offline browsing)
    js_content = (
        "/* Generated by sync.py on "
        + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%SZ")
        + " */\n"
        + f"window.GLOSSARY_DATA = {json.dumps(payload, indent=2, ensure_ascii=False)};\n"
    )
    DATA_JS.write_text(js_content, encoding="utf-8")

    m = payload["meta"]
    print(f"\n[Sync Complete]")
    print(f"  Wrote: {DATA_JSON.relative_to(HERE.parent)}")
    print(f"  Wrote: {DATA_JS.relative_to(HERE.parent)}")
    print(f"  Summary: {m['total_terms']} terms | {m['with_citations']} citations | {m['with_papers']} papers | {m['litigation_examples']} litigation examples")


# ═══════════════════════════════════════════════════════════════
# CLI & MAIN ENTRYPOINT
# ═══════════════════════════════════════════════════════════════
def main() -> None:
    parser = argparse.ArgumentParser(
        description="Sync glossary from SharePoint via Microsoft Graph API or local Excel."
    )
    parser.add_argument(
        "--local",
        nargs="?",
        const=str(DEFAULT_LOCAL_EXCEL),
        help="Sync from a local Excel file path instead of SharePoint.",
    )
    parser.add_argument(
        "--test-auth",
        action="store_true",
        help="Test Microsoft Graph API authentication and exit.",
    )
    parser.add_argument(
        "--setup",
        action="store_true",
        help="Display instructions for IT to register the Azure AD application.",
    )
    parser.add_argument(
        "file_arg",
        nargs="?",
        help="Optional local Excel file path (shorthand for --local <file>).",
    )

    args = parser.parse_args()
    load_dotenv()

    if args.setup:
        print_it_setup_instructions()
        return

    # Check for Azure AD credentials
    tenant_id = os.environ.get("AZURE_TENANT_ID")
    client_id = os.environ.get("AZURE_CLIENT_ID")
    client_secret = os.environ.get("AZURE_CLIENT_SECRET")
    site_url = os.environ.get("SHAREPOINT_SITE_URL")
    file_path = os.environ.get("SHAREPOINT_FILE_PATH")
    sharing_url = os.environ.get("SHAREPOINT_SHARING_URL")

    has_graph_creds = bool(tenant_id and client_id and client_secret)

    if args.test_auth:
        if not has_graph_creds:
            print("Error: Missing AZURE_TENANT_ID, AZURE_CLIENT_ID, or AZURE_CLIENT_SECRET in .env or environment.")
            sys.exit(1)
        client = GraphSyncClient(tenant_id, client_id, client_secret)
        print("[Graph API] Testing authentication to Azure AD...")
        token = client.get_token()
        print(f"[Graph API] Authentication successful! Acquired bearer token ({len(token)} chars).")
        return

    # Determine execution mode: Local vs. Microsoft Graph
    target_local = args.local or args.file_arg or os.environ.get("LOCAL_EXCEL_PATH")

    if target_local:
        excel_path = Path(target_local).expanduser().resolve()
        if not excel_path.exists():
            print(f"Error: Local Excel file not found: {excel_path}")
            sys.exit(1)
        print(f"[Local Mode] Reading from local file: {excel_path}")
        terms = read_excel(excel_path)
    else:
        # Microsoft Graph API mode
        if not has_graph_creds:
            print("\n[Notice] Azure AD credentials not found in environment or .env file.")
            print("To sync automatically from SharePoint, set up the Azure AD App credentials:")
            print_it_setup_instructions()
            print("Or run with a local Excel file: python sync.py --local path/to/file.xlsx\n")
            sys.exit(1)

        print("[Graph API Mode] Authenticating to Microsoft Graph API...")
        client = GraphSyncClient(tenant_id, client_id, client_secret)

        # Download file to a cache/temp directory
        cache_dir = HERE / ".cache"
        cache_dir.mkdir(parents=True, exist_ok=True)
        downloaded_excel = cache_dir / "sharepoint_glossary.xlsx"

        client.download_file(
            dest_path=downloaded_excel,
            site_url=site_url,
            file_path=file_path,
            sharing_url=sharing_url,
        )
        terms = read_excel(downloaded_excel)

    # Preserve rich metadata (papers, clusters, rings) from data.json
    enrich_with_existing_metadata(terms, DATA_JSON)

    payload = build_payload(terms)
    write_site_assets(payload)


if __name__ == "__main__":
    main()
